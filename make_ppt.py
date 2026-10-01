import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_COLOR = RGBColor(15, 23, 42)       # Slate 900
    CARD_BG = RGBColor(30, 41, 59)        # Slate 800
    CODE_BG = RGBColor(15, 17, 26)        # Deep Charcoal
    TEXT_MAIN = RGBColor(248, 250, 252)   # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)  # Slate 400
    ACCENT_BLUE = RGBColor(56, 189, 248)  # Sky 400
    ACCENT_GREEN = RGBColor(74, 222, 128) # Green 400
    ACCENT_ORANGE = RGBColor(251, 146, 60)# Orange 400
    ACCENT_RED = RGBColor(248, 113, 113)  # Red 400
    ACCENT_YELLOW = RGBColor(250, 204, 21)# Yellow 400

    FONT_TITLE = "Malgun Gothic"
    FONT_BODY = "Malgun Gothic"
    FONT_CODE = "Consolas"

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR

    def add_header(slide, main_category, sub_title):
        # Header Top Category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
        tf = cat_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = main_category
        p.font.name = FONT_TITLE
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE

        # Subtitle / Function Name
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.6))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = sub_title
        p_t.font.name = FONT_TITLE
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    slides_data = [
        {
            "cat": "트러블 슈팅 – 동적 메모리 관리 함수",
            "title": "malloc() – Memory Allocation",
            "diag_title": "📦 메모리 상태",
            "diag": "arr (0x2000)\n[ 0x5000 ] ----> 힙 메모리 (0x5000)\n                [ 쓰레기 | 쓰레기 | 쓰레기 | ... ] (40 byte)",
            "code": "int *arr = malloc(10 * sizeof(int));\n/* int 10개를 담을 공간을 할당 (40 byte)\n   지정한 크기만큼의 메모리 공간을 힙에 빌려줌 */",
            "bullets": [
                ("• 원하는 크기만큼 힙(Heap) 메모리 공간을 빌려오는 가장 기본적이고 빠른 함수임.", TEXT_MAIN, True),
                ("    + 빌려오기만 하면 끝이야? 컴퓨터가 알아서 다 준비해 줘?", ACCENT_ORANGE, False),
                ("• 공간만 빌려줄 뿐 청소를 안 해줘서, 이전에 쓰던 쓰레기 값이 그대로 들어있음.", TEXT_MUTED, False),
                ("    + 쓰레기 값이 있든 말든 내가 새 값 넣어서 덮어쓰면 되잖아?", ACCENT_ORANGE, False),
                ("• 값을 다 채우지 않은 상태에서 읽거나 포인터로 접근하면, 엉뚱한 쓰레기 주소를 찾아가 SIGSEGV 크래시가 발생함.", TEXT_MUTED, False)
            ],
            "footer": "// 아니;; 방을 빌려줬으면 최소한 기본 청소는 해줘야 하는 거 아니냐고..."
        },
        {
            "cat": "트러블 슈팅 – 동적 메모리 관리 함수",
            "title": "calloc() – Cleared Allocation",
            "diag_title": "📦 메모리 상태",
            "diag": "arr (0x2000)\n[ 0x6000 ] ----> 힙 메모리 (0x6000)\n                [  0  |  0  |  0  |  0  | ... ] (40 byte)",
            "code": "int *arr = calloc(10, sizeof(int));\n/* int 10개를 담을 공간을 할당 (40 byte)\n   메모리 안의 모든 비트를 0(NULL)으로 깨끗하게 정리 */",
            "bullets": [
                ("• 메모리 공간을 할당받자마자 모든 칸을 0(NULL)으로 깨끗하게 초기화해 줌.", TEXT_MAIN, True),
                ("    + malloc 쓰고 나서 내가 직접 for문 돌려서 0 넣는 거랑 뭐가 달라?", ACCENT_ORANGE, False),
                ("• 할당과 0 청소가 한 번에 이루어지므로, 초기화를 깜빡해서 생기는 버그를 사전에 완벽히 차단함.", TEXT_MUTED, False),
                ("    + 그럼 malloc 쓰지 말고 항상 calloc만 쓰면 되는 거 아냐?", ACCENT_ORANGE, False),
                ("• 모든 메모리를 0으로 미는 추가 작업이 들어가므로 malloc보다 미세하게 느림. 하지만 안전이 중요할 땐 calloc이 정답임.", TEXT_MUTED, False)
            ],
            "footer": "// 입주 청소까지 완벽하게 끝난 신축 원룸!"
        },
        {
            "cat": "트러블 슈팅 – 동적 메모리 관리 함수",
            "title": "realloc() – Re-Allocation",
            "diag_title": "📦 메모리 상태",
            "diag": "arr (0x2000)     --> [ 옛 10칸 방 (0x5000) ] (자동 폭파됨! 💥)\n                              │ (짐 몽땅 복사)\nnew_arr (0x2008) --> [ 20칸 새 방 (0x7000) ] (80 byte)",
            "code": "int *new_arr = realloc(arr, 20 * sizeof(int));\nif (new_arr != NULL) {\n    arr = new_arr; // 새 집 주소로 업데이트!\n}\n// arr의 메모리 공간을 80 byte로 변경함.",
            "bullets": [
                ("• 이미 빌려둔 메모리의 크기를 늘리거나 줄일 때 사용함.", TEXT_MAIN, True),
                ("    + 그냥 살던 방 벽 터서 제자리에서 늘리면 되는 거 아냐?", ACCENT_ORANGE, False),
                ("• 옆 방에 다른 데이터가 살고 있으면 제자리에서 못 늘려서, 짐을 몽땅 새 집으로 옮겨주고 예전 집은 폭파(free)해 버림.", TEXT_MUTED, False),
                ("    + 그럼 예전 집 주소(arr)는 어떻게 되는 건데?", ACCENT_ORANGE, False),
                ("• 예전 주소는 이미 사라진 폐허(Dangling Pointer)가 됨! 그래서 반드시 새 주소(new_arr)로 주소록을 업데이트해야 함.", TEXT_MUTED, False)
            ],
            "footer": "// 이사 갔으면 주소록 갱신해야지 옛날 주소 찾아가면 폭파된 폐허라고;;"
        },
        {
            "cat": "트러블 슈팅 – 동적 메모리 관리 함수",
            "title": "free() – Memory Free",
            "diag_title": "📦 메모리 상태",
            "diag": "arr (0x2000)\n[ NULL (0x0) ]\n[ 0x5000 방 반납 완료! (더 이상 접근 불가) ]",
            "code": "free(arr);\n// arr 메모리 공간을 운영체제에 반납.\narr = NULL;\n// arr을 실수로 다시 사용하는 것을 방지.",
            "bullets": [
                ("• 다 쓴 동적 메모리를 운영체제에 반납함.", TEXT_MAIN, True),
                ("    + 반납 안 하면 어떻게 되는데? 컴퓨터 끄면 어차피 사라지잖아?", ACCENT_ORANGE, False),
                ("• 프로그램이 실행되는 동안 반납하지 않은 메모리는 계속 쌓여서 결국 컴퓨터가 멈추는 '메모리 누수(Memory Leak)'가 일어남.", TEXT_MUTED, False),
                ("    + 그럼 free하고 나서 왜 굳이 arr = NULL을 해?", ACCENT_ORANGE, False),
                ("• 방을 뺐는데 주소록을 안 지우면, 이미 반납한 방을 또 반납(Double Free)하거나 찾아가서(Use-After-Free) 프로그램이 즉시 강제 종료됨.", TEXT_MUTED, False)
            ],
            "footer": "// 방 뺐으면 열쇠 반납하고 도어락 비밀번호도 지워야지!"
        },
        {
            "cat": "트러블 슈팅 – 메모리 조작 함수",
            "title": "memset() – Memory Set",
            "diag_title": "📦 메모리 상태",
            "diag": "arr[0..9] (스택 / 전역 / 힙 어디든 주소만 주면 OK!)\n[  0  |  0  |  0  |  0  |  0  |  0  | ... ] (40 byte 도배)",
            "code": "int arr[10]; // 스택(지역 변수) 배열\nmemset(arr, 0, sizeof(arr));\n// arr의 전체 크기(40 byte)만큼 0으로 싹 밀어버림.",
            "bullets": [
                ("• 이미 존재하는 특정 메모리 구간을 원하는 값(주로 0)으로 한 번에 도배함.", TEXT_MAIN, True),
                ("    + calloc도 0으로 채워주는데 memset이 왜 따로 필요해?", ACCENT_ORANGE, False),
                ("• calloc은 힙(동적 메모리) 전용이지만, memset은 스택(지역 변수), 전역 변수, 힙 가리지 않고 어디든 칠할 수 있음.", TEXT_MUTED, False),
                ("    + 그럼 memset으로 int 배열을 1로 다 채울 수도 있어?", ACCENT_ORANGE, False),
                ("• memset은 '바이트 단위'로 칠하기 때문에 int 배열에 1을 넣으면 0x01010101(16843009)이 들어가 버림! 0으로 밀 때만 안전함.", TEXT_MUTED, False)
            ],
            "footer": "// 어디든 0으로 도배해 주는 만능 페인트 롤러!"
        },
        {
            "cat": "트러블 슈팅 – 메모리 조작 함수",
            "title": "memcpy() – Memory Copy",
            "diag_title": "📦 메모리 상태",
            "diag": "original (0x5000): [ 10 | 20 | 30 | 40 ... ]\n                         │ 그대로 내용물 복사 (Deep Copy)\ncopy     (0x8000): [ 10 | 20 | 30 | 40 ... ] (완벽히 독립된 새 방!)",
            "code": "int *copy = malloc(10 * sizeof(int));\nmemcpy(copy, original, 10 * sizeof(int));\n// original의 내용물(40 byte)을 copy로 그대로 베껴 적음.",
            "bullets": [
                ("• 출발지(src)의 데이터를 목적지(dest)로 바이트 단위로 그대로 복사함.", TEXT_MAIN, True),
                ("    + 그냥 copy = original 이렇게 대입하면 복사 안 돼?", ACCENT_ORANGE, False),
                ("• 그렇게 하면 '주소'만 복사(얕은 복사)되어 둘이 같은 방을 쳐다보게 됨. original이 바뀌거나 free되면 copy도 같이 망가짐!", TEXT_MUTED, False),
                ("    + memcpy로 복사하면 뭐가 다른데?", ACCENT_ORANGE, False),
                ("• 새 방을 따로 빌려 내용물만 그대로 베껴 적는 '깊은 복사(Deep Copy)'이므로, 원본이 이사를 가든 폭파되든 내 복사본은 안전함.", TEXT_MUTED, False)
            ],
            "footer": "// 친구 일기장 주소만 적어갈래, 새 노트 사서 그대로 베껴 적어둘래?"
        }
    ]

    for item in slides_data:
        s = prs.slides.add_slide(blank_layout)
        set_slide_background(s)
        add_header(s, item["cat"], item["title"])

        # Left Card: Diagram & Code
        add_card(s, Inches(0.8), Inches(1.5), Inches(5.6), Inches(4.3), bg_color=CODE_BG, border_color=ACCENT_BLUE)
        c_box = s.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.2), Inches(4.1))
        tf_c = c_box.text_frame
        tf_c.word_wrap = True

        p = tf_c.paragraphs[0]
        p.text = item["diag_title"]
        p.font.name = FONT_TITLE
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GREEN
        p.space_after = Pt(6)

        p = tf_c.add_paragraph()
        p.text = item["diag"]
        p.font.name = FONT_CODE
        p.font.size = Pt(10.5)
        p.font.color.rgb = ACCENT_YELLOW
        p.space_after = Pt(12)

        p = tf_c.add_paragraph()
        p.text = "💻 예제 코드"
        p.font.name = FONT_TITLE
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        p.space_after = Pt(4)

        p = tf_c.add_paragraph()
        p.text = item["code"]
        p.font.name = FONT_CODE
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MAIN

        # Right Card: Conversational Q&A Bullets
        add_card(s, Inches(6.7), Inches(1.5), Inches(5.8), Inches(4.3))
        q_box = s.shapes.add_textbox(Inches(6.9), Inches(1.6), Inches(5.4), Inches(4.1))
        tf_q = q_box.text_frame
        tf_q.word_wrap = True

        for idx, (b_text, color, is_bold) in enumerate(item["bullets"]):
            p = tf_q.paragraphs[0] if idx == 0 else tf_q.add_paragraph()
            p.text = b_text
            p.font.name = FONT_BODY
            p.font.size = Pt(12)
            p.font.bold = is_bold
            p.font.color.rgb = color
            p.space_after = Pt(6)

        # Bottom Bar: Cheeky Footer Comment
        add_card(s, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.9), bg_color=CARD_BG, border_color=ACCENT_ORANGE)
        f_box = s.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.8))
        tf_f = f_box.text_frame
        tf_f.word_wrap = True
        p = tf_f.paragraphs[0]
        p.text = item["footer"]
        p.font.name = FONT_BODY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.italic = True
        p.font.color.rgb = ACCENT_RED

    out_path = r"C:\Coding\Jungle\5Week\5Week_Troubleshooting.pptx"
    prs.save(out_path)
    print(f"Presentation saved successfully to: {out_path}")

if __name__ == "__main__":
    create_deck()
